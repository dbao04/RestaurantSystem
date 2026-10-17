import docx
import os

doc_path = r"H:\NhaHang\NhaHang\BaoCao_KLTN_HeThongQuanLyNhaHang.docx"
doc = docx.Document(doc_path)

for i in range(910, 970):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        has_drawing = 'w:drawing' in p._element.xml
        print(f"P_{i} [has_img={has_drawing}]: {p.text.strip()}")

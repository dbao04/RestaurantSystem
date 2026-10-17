import docx
import os

doc_path = r"H:\NhaHang\NhaHang\BaoCao_KLTN_HeThongQuanLyNhaHang.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    xml = p._element.xml
    if "w:drawing" in xml or "w:pict" in xml:
        prev_t = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        curr_t = p.text.strip()
        next_t = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
        print(f"IMG [p_{i}]: curr='{curr_t}' | prev='{prev_t}' | next='{next_t}'")
